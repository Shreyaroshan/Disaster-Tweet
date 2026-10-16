"""Create the four-model final technical report from measured results.

Run with python-docx and Pillow. Codex uses its bundled artifact runtime for authoring.
"""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/final'
DATA=ROOT/'results/final_2026_10_16'


def main():
    records=json.loads((DATA/'comparison.json').read_text());by={r['model']:r for r in records}
    selections={name:json.loads((DATA/name/'selection.json').read_text()) for name in by}
    manifest=json.loads((ROOT/'data/processed/manifest.json').read_text())
    refs=json.loads((ROOT/'docs/references.json').read_text())
    uncertainty=json.loads((DATA/'paired_mlp_lstm.json').read_text())
    document=Document();document.core_properties.author='Disaster Tweet Project Team';document.core_properties.title='Comparative Analysis of Deep Learning Architectures for Disaster Tweet Classification';section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.85);section.bottom_margin=Inches(.85);section.left_margin=section.right_margin=Inches(1)
    section.different_first_page_header_footer=True
    for paragraph in section.first_page_footer.paragraphs:paragraph.clear()
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=document.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11 if name=='Normal' else 14 if name=='Heading 1' else 12)
        style.paragraph_format.space_after=Pt(7)
    for style in document.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    document.styles['Normal'].paragraph_format.line_spacing=1.12
    document.styles['Title'].font.size=Pt(21);document.styles['Title'].font.bold=True
    document.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    document.styles['Heading 2'].paragraph_format.space_before=Pt(6)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('Page ');r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    header=section.header.paragraphs[0];header.text='ICT 4442 Deep Learning Project';header.style='Caption'
    md=[]
    def p(text,bold=False):
        paragraph=document.add_paragraph();run=paragraph.add_run(text);run.bold=bold
        md.append(text+'\n');return paragraph
    def h(text,level=1):
        document.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
    def page():document.add_page_break()
    def table(headers,rows,weights=None):
        t=document.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        if weights is None:weights=[1]*len(headers)
        widths=[round(9360*w/sum(weights)) for w in weights];widths[-1]+=9360-sum(widths)
        tblpr=t._tbl.tblPr
        for tag in ['w:tblW','w:tblInd']:
            existing=tblpr.find(qn(tag))
            if existing is None:existing=OxmlElement(tag);tblpr.append(existing)
            existing.set(qn('w:type'),'dxa');existing.set(qn('w:w'),str(sum(widths) if tag=='w:tblW' else 0))
        for node in list(t._tbl.tblGrid):t._tbl.tblGrid.remove(node)
        for width in widths:
            node=OxmlElement('w:gridCol');node.set(qn('w:w'),str(width));t._tbl.tblGrid.append(node)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'CCCCCC');borders.append(node)
        tblpr.append(borders)
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();prevent=OxmlElement('w:cantSplit');trpr.append(prevent)
            if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr()
                width=tcpr.find(qn('w:tcW'));width.set(qn('w:w'),str(widths[j]));width.set(qn('w:type'),'dxa')
                margins=OxmlElement('w:tcMar')
                for edge,value in [('top',70),('bottom',70),('start',90),('end',90)]:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');margins.append(node)
                tcpr.append(margins)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'333333' if i==0 else 'F5F5F5' if i%2==0 else 'FFFFFF');tcpr.append(shade)
                for paragraph in c.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.space_before=Pt(2);paragraph.paragraph_format.line_spacing=1.05
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(c.text))>35 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(9.5);run.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
        md.append('| '+' | '.join(headers)+' |\n|'+ '|'.join(['---']*len(headers))+'|')
        for row in rows:md.append('| '+' | '.join(str(v).replace('\n',' ') for v in row)+' |')
        md.append('');document.add_paragraph().paragraph_format.space_after=Pt(1)
        return t
    def figure(filename,caption,width=6.2):
        paragraph=document.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(OUT/'figures'/filename),width=Inches(width))
        paragraph.paragraph_format.space_after=Pt(3)
        cap=document.add_paragraph(caption,style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        md.append(f'![{caption}](figures/{filename})\n')
    # Cover
    title=document.add_paragraph('Comparative Analysis of Deep Learning Architectures for Disaster Tweet Classification',style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for text in ['Final Technical Report','ICT 4442 Deep Learning','School of Computer Engineering','Manipal Institute of Technology','Manipal Academy of Higher Education']:
        para=document.add_paragraph(text);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    for text in ['Shreya Roshan   230911046','Tanisha Mathur   230953110','Adithya Kini   230911008','Arnab Sinha   230911356']:
        para=document.add_paragraph(text);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    for text in ['October 2026','Technical implementation and measured evaluation completed','Team methodology review and contribution signatures pending']:
        para=document.add_paragraph(text);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    md.append('# Comparative Analysis of Deep Learning Architectures for Disaster Tweet Classification\n')
    page();h('Abstract')
    p(f"This project compares a multilayer perceptron, a one-dimensional convolutional neural network, a bidirectional long short-term memory network, and an encoder-only Transformer for binary disaster-tweet classification. All models use the same text-only cleaning pipeline, training-fitted vocabulary, GloVe word-vector source, and saved train-validation-test partition. After excluding conflicting normalized labels and repeated normalized tweets, 6,815 unique labeled examples remain. Models are selected on validation data and evaluated on the same 1,023-row local test set. MLP achieves the highest observed accuracy at {by['mlp']['accuracy']*100:.2f}%, while LSTM has the highest observed F1 at {by['lstm']['f1']*100:.2f}% and recall at {by['lstm']['recall']*100:.2f}%. The F1 difference between these models is small. The results show a tradeoff between conservative predictions and disaster recall rather than a universal benefit from architectural complexity. Reused test data, label ambiguity, different historical search budgets, and the absence of independent external evaluation limit generalization claims.")
    p('Keywords: disaster tweets; text classification; MLP; CNN; bidirectional LSTM; encoder-only Transformer; GloVe.',bold=True)
    h('1 Introduction')
    p('Social-media posts can contain useful descriptions of emergencies, but short messages also contain jokes, figurative expressions, references to past events and missing context. A word such as fire can describe a hazard or enthusiastic slang. This project classifies the supplied tweet text according to the dataset label; it does not verify whether an event actually happened or whether a user is trustworthy.')
    p('The research question is how feed-forward, convolutional, recurrent and self-attention architectures differ under common data preparation and evaluation. The study emphasizes reproducible comparison, saved experimental evidence and qualitative error analysis rather than a production emergency-response system.')
    h('2 Objectives')
    for text in ['Implement four distinct architecture families using a shared preprocessing and evaluation pipeline.','Compare accuracy, precision, recall, positive-class F1, ROC AUC, parameter count and measured CPU training cost.','Analyze false positives and false negatives and identify limitations involving ambiguity, context and data quality.']:p(text)
    page();h('3 Literature Review')
    p('Disaster-classification studies [1–5] investigate contextual models, CNNs, cross-event transfer and operational categories. Their datasets, labels, auxiliary features and pretraining differ, so published numbers provide context rather than directly comparable target scores.')
    p('The four families are grounded in unordered composition [7], sentence CNNs [6,9], recurrent memory [10] and self-attention [11]. All use common GloVe vectors [12] and AdamW [13]. Unlike BERT [8], the encoder here has no contextual language-model pretraining. This distinction separates representation resources from architecture.')
    h('3 1 Comparative literature summary',2)
    labels=['Balakrishnan 2022','Dharrao 2024','Nguyen 2017','Kabir 2019','HumAID 2021','Kim 2014','Iyyer 2015','BERT 2019','Zhang 2015','LSTM 1997','Transformer 2017','GloVe 2014','AdamW 2019']
    short=['Contextual classifiers; different split and pretraining.','BERT and CNN; a different evaluation protocol.','Crisis transfer; motivates event-held-out testing.','Rescue scheduling extends beyond binary text labels.','Larger humanitarian multiclass benchmark.','Motivates local phrase filters and pooling.','Motivates a simple unordered neural baseline.','Contextual pretraining is a separate condition.','Supports validation-based hyperparameter search.','Memory gates motivate the recurrent model.','Self-attention and positions motivate the encoder.','Supplies the shared static word vectors.','Supplies the common optimization mechanism.']
    table(['Study','Method or task','Relevance to this comparison'],[[f"[{r['id']}] {labels[i]}",r['method'],short[i]] for i,r in enumerate(refs)],[1.5,2,3])
    p('These studies motivate architectural diversity but do not identify a universal winner. The common data protocol and selected hyperparameters are detailed below.')
    page();h('4 Dataset and Shared Preprocessing')
    p('The primary dataset is Kaggle Natural Language Processing with Disaster Tweets [14]. Its labeled training CSV has 7,613 rows; the separate competition test CSV has 3,263 rows without labels. Files were acquired through a pinned public mirror, and source URLs plus SHA-256 checksums are recorded. Only text is an input feature. ID identifies rows, target supplies the label, and keyword/location are excluded.')
    table(['Processing stage','Rows or count'],[['Original labeled rows','7,613'],['Conflicting-label or empty rows excluded','282'],['Conflicting normalized texts','79'],['Repeated same-label normalized rows removed','516'],['Retained unique labeled rows','6,815'],['Training vocabulary including PAD and UNK','12,089']],[4.5,2])
    p('The single preprocessing module decodes HTML, lowercases text, replaces URLs with urltoken and mentions with usertoken, and extracts alphanumeric words with internal apostrophes. Hashtag words and negation are retained. No stopword deletion or stemming is applied. For each normalized text with consistent labels, one row is retained; normalized texts with contradictory labels are excluded before splitting. This global label-consistency filtering changes the benchmark population and must be disclosed.')
    table(['Partition','Rows','Disaster','Non disaster'],[[name.title(),str(info['rows']),str(info['positive']),str(info['rows']-info['positive'])] for name,info in manifest['splits'].items()],[2,1.5,1.5,1.5])
    p('A stratified 70/15/15 split with seed 42 is saved and reused. The vocabulary is fitted on training text only, capped at 20,000 entries, with PAD=0 and UNK=1. Encodings contain up to 64 tokens with right padding; an empty cleaned input becomes UNK. Split IDs, file checksums and vocabulary checksum are validated before training. Identical normalized texts cannot cross partitions, although different tweets about the same event may still do so.')
    h('4 1 Common word representations',2)
    p('All models start from GloVe 6B 100-dimensional vectors [12], using a checksum-verified Gensim-format mirror. The source covers 10,227 of the 12,087 training vocabulary words, or 89.14% of training token occurrences. Missing training words receive deterministic small random vectors, UNK uses the mean matched vector, and PAD is zero. Model-specific fine-tuning or freezing is a selected hyperparameter, not a separate source of unlabeled pretraining.')
    page();h('5 Methodology')
    h('5 1 Multilayer perceptron assigned to Member 2',2)
    p('The MLP maps token IDs to 100-dimensional embeddings, then concatenates masked mean and masked maximum pooling into a 200-dimensional vector. Layer normalization and dropout precede a 128-unit dense ReLU layer, a second dropout layer and a single output logit. The embeddings are fine-tuned. Pooling is invariant to added right padding and discards token order, so the model offers a simple baseline for whether global lexical evidence is sufficient.')
    p('The selected model uses dropout 0.5 and a validation-selected threshold of 0.49. Its probability is sigmoid(logit); binary cross-entropy with logits is used during training. The retained September configuration was reproduced under the new common trainer at three seeds. The final validation selection retained the designated seed-42 model. Historical larger-embedding, TF-IDF and mixture experiments remain archived rather than being presented as improvements they did not establish.')
    h('5 2 One dimensional CNN assigned to Member 1',2)
    p('The CNN applies parallel one-dimensional convolutions of widths 2, 3 and 4 over the word embeddings, with 64 filters per width. ReLU activations and masked global max pooling produce a 192-dimensional phrase vector. Embedding-channel dropout is 0.1, classifier dropout is 0.5, and a linear layer produces the output logit. Embeddings are fine-tuned and the selected threshold is 0.50.')
    p('Padding-only windows are excluded from pooling. Tweets shorter than a kernel retain one zero-padded window; batch right-padding is trimmed to avoid unnecessary computation. Convolutions can detect local lexical patterns and preserve more short-range order than the MLP, but global max pooling retains only the strongest activation for each filter. This implementation follows the general sentence-CNN idea [6] without claiming to reproduce its benchmark results.')
    h('5 3 Common training and inference behavior',2)
    p('All model classes return a scalar logit per example. Shared code controls data loading, optimizer setup, deterministic seeds, minibatches, gradient clipping, checkpoint selection, threshold selection, probability generation and output files. Each owner maintains only their model class, selected configuration and methodology card. This avoids separate cleaning implementations drifting across member branches.')
    page();h('5 4 Bidirectional LSTM assigned to Member 4',2)
    p('The selected recurrent model uses one bidirectional LSTM layer with 64 hidden units in each direction. Fine-tuned GloVe embeddings are passed through input dropout 0.15. Length-aware packed sequences prevent padding from participating in recurrent updates. Masked mean and max pooling over the 128-dimensional recurrent outputs are concatenated, normalized and classified by a 64-unit ReLU head with dropout 0.4.')
    p('The forward and backward recurrent states provide context from both directions. Gating is motivated by the LSTM memory mechanism [10], but the implementation is a modern bidirectional library layer rather than a reproduction of the original synthetic-task experiments. Four candidate configurations varied hidden size, layer count and embedding freezing. The top two were repeated at two further seeds; the final configuration was selected by mean validation accuracy. Validation chose one seed-42 model at threshold 0.48.')
    h('5 5 Encoder only Transformer assigned to Member 3',2)
    p('The selected Transformer projects the common 100-dimensional embeddings to 64 dimensions and adds sinusoidal positional encoding. One encoder block uses four self-attention heads and a 128-dimensional feed-forward sublayer, with GELU, residual connections, normalization and dropout 0.3. A key-padding mask excludes padded positions from attention. Masked mean and max pooling feed a 64-unit GELU classifier and one logit.')
    p('Attention is bidirectional across the available tweet tokens. There is no causal mask, decoder, text generation or pretrained BERT encoder. The architecture draws on the encoder mechanism in [11], while all attention/classifier weights are trained on this dataset. The selected condition freezes the GloVe table. Four candidates varied width, depth and freezing; top candidates were repeated across seeds, and validation selected threshold 0.47 for one seed-42 model. Distinct encoder layers are initialized independently in deeper candidates.')
    h('5 6 Ownership and review status',2)
    p('Assignments follow the latest team instruction: Member 1 CNN, Member 2 MLP, Member 3 encoder-only Transformer, Member 4 LSTM. This supersedes the synopsis mapping. The name-to-member-number mapping must be confirmed by the team. These model descriptions are AI-assisted drafts for each owner to review, adapt and explain; they are not signed claims of individual implementation authorship.')
    page();h('6 Experimental Setup')
    p('Experiments ran on an ARM64 CPU with four PyTorch threads, deterministic algorithms and Python 3.13.1. Exact tested package versions and source/data checksums accompany the repository. AdamW [13] optimizes unweighted binary cross-entropy, with batch size 64, weight decay 0.01 and gradient norm clipping at 5. Training embeddings use learning rate 0.0002; dense/recurrent/CNN layers use 0.001 and Transformer layers 0.0005.')
    table(['Model','Epoch cap','Patience','Dropout','GloVe'],[['MLP','40','7','0.5','Fine tuned'],['CNN','30','6','0.5','Fine tuned'],['LSTM','30','6','0.4','Fine tuned'],['Transformer','30','6','0.3','Frozen']],[1.6,1.1,1.1,1.1,1.6])
    p('The operational checkpoint objective is validation accuracy, with positive-class F1 breaking ties, reflecting the team’s later emphasis on accuracy. The synopsis originally emphasized F1; the final analysis therefore reports and discusses both metrics explicitly. MLP and CNN each reproduce one previously selected configuration across seeds 42, 123 and 2026. Each new family screens four configurations at seed 42, then repeats its top two configurations at seeds 123 and 2026. The winner is chosen by mean validation accuracy across those seeds.')
    p('For each winning configuration, validation compares the predesignated seed-42 model and the equal-weight three-seed ensemble over thresholds 0.35 to 0.65 in 0.01 steps. Accuracy, F1, closeness to 0.50 and preference for a single model break ties in that order. All four selected predictors are single models. A final lock file records every selection hash before test labels are loaded for the integrated comparison.')
    p('The current integrated stage contains 22 training fits: 3 MLP, 3 CNN, 8 LSTM and 8 Transformer. This excludes substantial historical MLP/CNN searches, which are archived and have unequal budgets. Thus common preprocessing and metrics do not make this a controlled comparison of architecture alone. Search cost and selected-fit cost are reported separately; CPU timings are measurements from this machine, not universal speed claims.')
    h('6 1 Evaluation definitions',2)
    p('Accuracy is the fraction of all correct labels. Precision is the fraction of predicted disaster tweets labeled disaster. Recall is the fraction of labeled disasters detected. F1 is the harmonic mean of precision and recall for the disaster class, not macro F1. ROC AUC summarizes ranking across thresholds. Confusion counts show the practical false-alarm and missed-disaster tradeoff. The unlabeled competition test is used only to export prediction files, not to compute these scores.')
    page();h('7 Results and Discussion')
    table(['Model','Accuracy %','Precision %','Recall %','F1 %','ROC AUC'],[[r['model'].upper() if r['model']!='transformer' else 'Transformer',f"{r['accuracy']*100:.2f}",f"{r['precision']*100:.2f}",f"{r['recall']*100:.2f}",f"{r['f1']*100:.2f}",f"{r['roc_auc']:.4f}"] for r in records],[1.5,1,1,1,1,1])
    figure('comparison.png','Figure 1 Common local test metrics for the four selected models')
    p('MLP has the highest observed accuracy and precision. LSTM has the highest observed recall, F1 and ROC AUC. LSTM detects 17 more disasters than MLP, but produces 25 more false alarms, leaving eight fewer total correct labels. This illustrates why the preferred model depends on the relative importance of missed events and false alarms. The observed F1 difference is only about 0.22 percentage points and does not establish statistical superiority.')
    p('The CNN detects local phrase evidence but generates more false positives than MLP. The encoder-only Transformer is comparatively conservative at its selected threshold and misses more labeled disasters. Its attention layers have no large-scale language-model pretraining, so its result should not be interpreted as evidence against pretrained contextual models such as BERT [8].')
    page();h('7 1 Confusion matrices')
    figure('confusion_matrices.png','Figure 2 Confusion counts with class 0 non disaster and class 1 disaster',width=5.8)
    table(['Model','True disaster','Missed disaster','False alarm','Correct non disaster'],[[r['model'].upper(),str(r['true_positive']),str(r['false_negative']),str(r['false_positive']),str(r['true_negative'])] for r in records],[1.5,1.2,1.2,1.2,1.4])
    page();h('7 2 Error analysis')
    p('All errors are saved with raw text, label, probability and prediction. A deterministic review sample includes high-confidence false positives and false negatives from each family. The examples below illustrate ambiguity and possible annotation/context limitations; no label was manually corrected. Candidate interpretations require member review and should not be treated as certified relabeling.')
    table(['Tweet ID and example','Dataset label and observed error','Interpretation for team review'],[
        ['8489  I JUST SCREAMED','Label 1; MLP and LSTM predict 0','The text resembles enthusiasm and lacks a literal incident. Missing URL/context or an ambiguous label may matter.'],
        ['467  Armageddon plus a URL','Label 1; MLP predicts 0','One word plus an inaccessible link offers insufficient text context for event classification.'],
        ['6211  Parole for a school bus hijacker','Label 0; CNN and LSTM predict 1','Crime-related wording can trigger hazard features even when the story concerns a past event or legal decision.'],
        ['10090  Satellite spies a super typhoon','Label 0; Transformer predicts 1','Strong disaster vocabulary conflicts with the negative annotation; the task boundary or label/context needs review.']],[1.8,1.7,3])
    p('Model behavior and annotation ambiguity can coexist. MLP pooling loses word order and can miss context; CNN max pooling can amplify a hazardous phrase without time context; recurrent state can improve recall but still confuse old events and literal disasters; attention cannot recover information absent from the supplied text. Normalized duplicate removal prevents exact text leakage, but it does not solve semantic label inconsistency or event overlap.')
    h('7 3 Computational comparison',2)
    table(['Model','Parameters','Selected fit seconds','This stage search seconds','Fits'],[[r['model'].upper(),f"{r['total_parameters']:,}",f"{r['training_seconds_selected_fits']:.2f}",f"{r['training_seconds_search']:.2f}",str(r['search_fits'])] for r in records],[1.5,1.5,1.3,1.5,.7])
    p('The Transformer has 1,257,413 total parameters, but only 48,513 trainable parameters because its embedding table is frozen. Other selected models fine-tune their word vectors. The parameter table therefore reports model storage/capacity separately from trainability. Inference timings include loading and vary with operating-system caches; exact per-model observations remain in the metrics files.')
    page();h('7 4 Uncertainty and limitations')
    intervals=uncertainty['95_percent_intervals']
    p(f"A 2,000-sample paired tweet bootstrap compares LSTM minus MLP. The conditional 95% interval for accuracy change is {intervals['accuracy'][0]*100:+.2f} to {intervals['accuracy'][1]*100:+.2f} percentage points; for F1 it is {intervals['f1'][0]*100:+.2f} to {intervals['f1'][1]*100:+.2f} points. Both include zero. The small observed difference should therefore be described as exploratory rather than statistically established.")
    p('The local test split was already evaluated during earlier project iterations, including unsuccessful MLP follow-ups. Although this integrated evaluation locks selections before reading test labels, it is not a fresh independent test. Validation is also reused for early stopping, configuration selection and threshold selection. The bootstrap conditions on fitted predictors and ignores training/search uncertainty and correlations among tweets about the same event.')
    p('Additional limitations include the small supervised dataset, random rather than event-held-out partitioning, global filtering of contradictory normalized labels, possible non-exact duplicate overlap, ambiguous annotations, excluded metadata, and one language/domain. Different search budgets, freezing choices and historical tuning mean differences cannot be attributed solely to architectural family. Larger neural or pretrained contextual models may behave differently under more data and independent evaluation.')
    h('8 Conclusion')
    p(f"The project implements and measures all four proposed architecture families under common data preparation and output formats. MLP achieves {by['mlp']['accuracy']*100:.2f}% accuracy and LSTM achieves {by['lstm']['f1']*100:.2f}% positive-class F1 on the reused local test split. A simple pooled representation remains competitive, while recurrent context yields higher disaster recall. CNN and self-attention provide alternative inductive biases but do not uniformly improve every metric. The measured differences support a discussion of tradeoffs and limitations rather than a universal winning architecture.")
    h('9 Future Work')
    p('Obtain a fresh independently labeled evaluation set or event-held-out split; audit ambiguous labels with multiple annotators; compare pretrained contextual encoders as an explicitly separate transfer-learning condition; calibrate probabilities and choose thresholds from application costs; investigate event/time/domain shifts; and extend to humanitarian categories such as HumAID [5] with an appropriate multiclass protocol.')
    h('10 Acknowledgement and AI Assistance')
    p('The team acknowledges the dataset providers, GloVe authors, open-source framework maintainers and cited researchers. OpenAI Codex assisted with dataset acquisition, repository restructuring, code, model training, tests, literature synthesis, figures and this report draft. Results are actual local measurements, not literature values or fabricated scores. Each member must verify their assigned model and methodology, document actual review/reproduction/changes, and complete the signed contribution page. Required institutional plagiarism/AI-content checks have not been generated or certified here.')
    page();h('References')
    for ref in refs:
        paragraph=p(f"[{ref['id']}] {ref['authors']}, {ref['title']}, {ref['venue']}, {ref['year']}. {ref['url']}")
        paragraph.paragraph_format.space_after=Pt(6)
        for run in paragraph.runs:run.font.size=Pt(10)
    p('[14] Kaggle, Natural Language Processing with Disaster Tweets, dataset and competition. https://www.kaggle.com/competitions/nlp-getting-started/data')
    page();h('Individual Contribution Statement')
    p('Mandatory extra page for team verification and signatures',bold=True)
    p('Technical implementation and measurements were initially generated with AI assistance as disclosed above. The table records intended model responsibility; it does not certify authorship. Each member should enter their actual review, reproduction, code changes, writing and commit/PR evidence before signing. Merely pushing generated files is not an implementation contribution claim.')
    table(['Assigned member','Model responsibility','Verified work and commit evidence'],[['Member 1 user','CNN','To be completed by member'],['Member 2','MLP','To be completed by member'],['Member 3','Encoder only Transformer','To be completed by member'],['Member 4','LSTM','To be completed by member']],[1.4,1.7,3.4])
    p('The team should confirm the name-to-member mapping below. No signature or contribution has been filled on behalf of a member.')
    for name,roll in [('Shreya Roshan','230911046'),('Tanisha Mathur','230953110'),('Adithya Kini','230911008'),('Arnab Sinha','230911356')]:
        p(f'{name}   {roll}');p('Assigned model __________________  Signature __________________  Date __________')
    p('Before submission: apply the official Part C template if required, add the team number, confirm ownership, complete signatures, and obtain the institutional checks. Only the guidelines and synopsis were available in references; the separate Part C template was not supplied.')
    OUT.mkdir(parents=True,exist_ok=True);document.save(OUT/'Final_Report_Disaster_Tweet_Classification.docx')
    (OUT/'final_report.md').write_text('\n'.join(md),encoding='utf-8')
    print(OUT/'Final_Report_Disaster_Tweet_Classification.docx')


if __name__=='__main__':main()
